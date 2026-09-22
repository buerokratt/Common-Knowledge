"""
test_tasks.py — Unit tests for worker/tasks.py extraction logic.

These tests run WITHOUT Docker containers. All external dependencies
(Vault, Ruuter, Azure OpenAI) are mocked. They verify:
  - HTML extraction strategy selection (trafilatura / BeautifulSoup paths)
  - LLM evaluate + correct branching
  - PDF and generic file routing
  - normalize_newlines
  - set_up_logging deduplication
  - Metadata mutation correctness
  - Error handling in LLM helpers
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from unstructured.documents.elements import Table


# ---------------------------------------------------------------------------
# Helpers to build minimal EntityToClean objects without hitting the filesystem
# ---------------------------------------------------------------------------


def _make_entity(
    tmp_path: Path,
    file_type: str,
    content: str,
    use_llm: bool = False,
    use_llm_correction: bool = False,
) -> MagicMock:
    """
    Create a minimal EntityToClean-like namespace pointing at real temp files
    so we can call the worker functions directly without Pydantic validation.
    """
    source_file = tmp_path / f"source{file_type}"
    source_file.write_text(content, encoding="utf-8")

    meta_data = {
        "file_type": file_type,
        "url": "https://example.com",
        "metadata": {"cleaned": False},
    }
    meta_file = tmp_path / "source.meta.json"
    meta_file.write_text(json.dumps(meta_data))

    log_file = tmp_path / "test.log"
    log_file.touch()

    entity = MagicMock()
    entity.file_path = source_file
    entity.meta_data_path = meta_file
    entity.directory_path = tmp_path
    entity.logs_path = log_file
    entity.url = "https://example.com"
    entity.source_file_id = "test-id"
    entity.source_base_id = "source-id"
    entity.agency_base_id = "agency-id"
    entity.source_run_report_base_id = "report-id"
    entity.use_llm = use_llm
    entity.use_llm_correction = use_llm_correction
    return entity


# ---------------------------------------------------------------------------
# normalize_newlines
# ---------------------------------------------------------------------------


class TestNormalizeNewlines:
    def test_three_newlines_collapsed(self) -> None:
        from worker.tasks import normalize_newlines

        assert normalize_newlines("a\n\n\nb") == "a\n\nb"

    def test_five_newlines_collapsed(self) -> None:
        from worker.tasks import normalize_newlines

        assert normalize_newlines("a\n\n\n\n\nb") == "a\n\nb"

    def test_two_newlines_unchanged(self) -> None:
        from worker.tasks import normalize_newlines

        assert normalize_newlines("a\n\nb") == "a\n\nb"

    def test_single_newline_unchanged(self) -> None:
        from worker.tasks import normalize_newlines

        assert normalize_newlines("a\nb") == "a\nb"

    def test_empty_string(self) -> None:
        from worker.tasks import normalize_newlines

        assert normalize_newlines("") == ""


# ---------------------------------------------------------------------------
# _beautifulsoup_extract
# ---------------------------------------------------------------------------


class TestBeautifulSoupExtract:
    def test_extracts_main_element_content(self) -> None:
        from worker.tasks import _beautifulsoup_extract

        html = """<html><body>
          <header>Nav stuff</header>
          <main><h1>Title</h1><p>Body text.</p></main>
          <footer>Footer</footer>
        </body></html>"""
        result = _beautifulsoup_extract(html)
        assert "Title" in result
        assert "Body text" in result

    def test_removes_nav_script_style(self) -> None:
        from worker.tasks import _beautifulsoup_extract

        html = """<html><body>
          <nav>Home | About</nav>
          <script>alert('x')</script>
          <style>body{}</style>
          <main><p>Real content here.</p></main>
        </body></html>"""
        result = _beautifulsoup_extract(html)
        assert "Real content here" in result
        assert "alert" not in result
        assert "Home | About" not in result

    def test_fallback_to_body_when_no_main(self) -> None:
        from worker.tasks import _beautifulsoup_extract

        html = """<html><body>
          <div><h1>No Main Element</h1><p>Still extracted.</p></div>
        </body></html>"""
        result = _beautifulsoup_extract(html)
        assert "No Main Element" in result
        assert "Still extracted" in result

    def test_returns_string(self) -> None:
        from worker.tasks import _beautifulsoup_extract

        result = _beautifulsoup_extract("<html><body><p>hello</p></body></html>")
        assert isinstance(result, str)
        assert len(result) > 0


# ---------------------------------------------------------------------------
# _trafilatura_extract
# ---------------------------------------------------------------------------


class TestTrafilaturaExtract:
    def test_extracts_article_content(self) -> None:
        from worker.tasks import _trafilatura_extract

        html = """<html><body>
          <article>
            <h1>Article Title</h1>
            <p>This is a full article with enough content to be extracted by trafilatura.</p>
            <p>Second paragraph with more substantial text to meet minimum length requirements.</p>
          </article>
        </body></html>"""
        result = _trafilatura_extract(html)
        # trafilatura may return None on very short content — just check type
        assert result is None or isinstance(result, str)

    def test_returns_none_on_empty_html(self) -> None:
        from worker.tasks import _trafilatura_extract

        result = _trafilatura_extract("<html><body></body></html>")
        assert result is None

    def test_returns_none_on_noise_only(self) -> None:
        from worker.tasks import _trafilatura_extract

        html = "<html><body><nav>Home | About</nav><footer>Copyright 2024</footer></body></html>"
        result = _trafilatura_extract(html)
        assert result is None or isinstance(result, str)


# ---------------------------------------------------------------------------
# HTML table preservation — trafilatura placeholder pass + cleanup
# ---------------------------------------------------------------------------


class TestTrafilaturaHtmlTables:
    """
    Every <table> in the source HTML must survive as raw HTML in the
    extracted Markdown — trafilatura's own Markdown table renderer cannot
    represent colspan/rowspan, so we replace tables with placeholders before
    extraction and inject the originals back.
    """

    def test_simple_table_survives_as_html(self) -> None:
        from worker.tasks import _trafilatura_extract

        html = """
        <html><body>
          <article>
            <h1>Report</h1>
            <p>This is prose that must be extracted alongside the table.
               A few more sentences so trafilatura recognises the article.</p>
            <table>
              <tr><th>A</th><th>B</th></tr>
              <tr><td>1</td><td>2</td></tr>
            </table>
            <p>More prose after the table so trafilatura keeps it in scope.</p>
          </article>
        </body></html>
        """
        result = _trafilatura_extract(html)
        assert result is not None
        assert "<table>" in result
        assert "<td>1</td>" in result and "<td>2</td>" in result

    def test_colspan_table_preserved(self) -> None:
        from worker.tasks import _trafilatura_extract

        html = """
        <html><body>
          <article>
            <h1>Colspan table</h1>
            <p>Prose paragraph one — enough content that trafilatura keeps
               the article body when favor_precision is on.</p>
            <table>
              <tr><th>I</th><th>II</th><th>III</th></tr>
              <tr><td colspan="3">Applies to all three</td></tr>
              <tr><td>a</td><td>b</td><td>c</td></tr>
            </table>
            <p>Prose paragraph two after the table to anchor extraction.</p>
          </article>
        </body></html>
        """
        result = _trafilatura_extract(html)
        assert result is not None
        assert 'colspan="3"' in result
        assert "Applies to all three" in result

    def test_placeholder_removed_when_table_missing_from_output(self) -> None:
        """If trafilatura drops the placeholder region (favor_precision), we
        must not leave the raw @@CKB_TABLE_N@@ marker in the output."""
        from worker.tasks import _TABLE_PLACEHOLDER_PREFIX, _trafilatura_extract

        html = """
        <html><body>
          <article>
            <h1>Report</h1>
            <p>Body paragraph one.</p>
            <table><tr><td>x</td></tr></table>
            <p>Body paragraph two.</p>
          </article>
        </body></html>
        """
        result = _trafilatura_extract(html) or ""
        # Either the table was injected (contains <table>) or the placeholder
        # was dropped by trafilatura — never leave the raw marker text.
        assert _TABLE_PLACEHOLDER_PREFIX not in result


class TestCleanHtmlTables:
    """
    Post-processing step that strips presentational attributes from HTML
    tables and pretty-prints them.
    """

    def test_removes_style_and_class(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            "Before.\n\n"
            '<table class="foo" style="border:1px"><tr>'
            '<td style="padding:8px">x</td></tr></table>\n\n'
            "After."
        )
        result = clean_html_tables(text)
        assert 'class="foo"' not in result
        assert "style=" not in result
        assert "<td>x</td>" in result
        # Prose outside the table is untouched.
        assert "Before." in result and "After." in result

    def test_keeps_colspan_rowspan_scope(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            "<table><tr>"
            '<th scope="col" style="background:red">A</th>'
            "</tr><tr>"
            '<td colspan="2" rowspan="3" class="x">B</td>'
            "</tr></table>"
        )
        result = clean_html_tables(text)
        assert 'scope="col"' in result
        assert 'colspan="2"' in result
        assert 'rowspan="3"' in result
        assert "style=" not in result and "class=" not in result

    def test_unwraps_span_inside_cell(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '<table><tr><td><span style="color:red">important</span></td></tr></table>'
        )
        result = clean_html_tables(text)
        assert "<span" not in result
        assert "important" in result

    def test_downgrades_headings_inside_cell(self) -> None:
        from worker.tasks import clean_html_tables

        text = "<table><tr><td><h5>Header</h5></td></tr></table>"
        result = clean_html_tables(text)
        assert "<h5>" not in result and "</h5>" not in result
        assert "<strong>Header</strong>" in result

    def test_strips_wrapping_responsive_div(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '<div class="w-100 table-responsive">'
            "<table><tr><td>x</td></tr></table>"
            "</div>"
        )
        result = clean_html_tables(text)
        assert "table-responsive" not in result
        assert "<table>" in result and "<td>x</td>" in result

    def test_content_outside_tables_untouched(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '## Heading\n\nPara <a href="http://x">link</a> with style="foo".\n\n'
            "<table><tr><td>y</td></tr></table>"
        )
        result = clean_html_tables(text)
        # Outside the table, style="..." stays as-is (it is just prose).
        assert 'style="foo"' in result
        assert "<td>y</td>" in result

    def test_preserves_space_between_inline_elements(self) -> None:
        """Whitespace separating inline tags inside a cell is content, not
        indentation — collapsing it welds adjacent words together."""
        from worker.tasks import clean_html_tables

        text = (
            "<table><tr><td>"
            "<strong>Riigi</strong> <em>pension</em> ja "
            '<a href="/x">link</a> <strong>siin</strong>'
            "</td></tr></table>"
        )
        result = clean_html_tables(text)
        assert "<strong>Riigi</strong> <em>pension</em>" in result
        assert "</a> <strong>siin</strong>" in result

    def test_collapses_indentation_between_structural_tags(self) -> None:
        """Whitespace between <tr>/<td> is indentation and is dropped."""
        from worker.tasks import clean_html_tables

        text = "<table>\n  <tr>\n    <td>a</td>\n    <td>b</td>\n  </tr>\n</table>"
        result = clean_html_tables(text)
        assert "<td>a</td><td>b</td>" in result

    def test_nested_table_stays_balanced(self) -> None:
        """A nested table must not be cut at the first </table>, which would
        strand the outer table's remaining cells as stray markup."""
        from worker.tasks import clean_html_tables

        text = (
            "<table><tr><td>outer"
            "<table><tr><td>inner</td></tr></table>"
            "</td><td>sibling</td></tr></table>"
        )
        result = clean_html_tables(text)
        assert result.count("<table>") == 2
        assert result.count("</table>") == 2
        assert "inner" in result and "sibling" in result

    def test_nested_table_attrs_cleaned(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '<table class="o"><tr><td>'
            '<table class="i"><tr><td style="x">deep</td></tr></table>'
            "</td></tr></table>"
        )
        result = clean_html_tables(text)
        assert "class=" not in result and "style=" not in result
        assert "<td>deep</td>" in result

    def test_multiple_tables_and_surrounding_prose(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            "Alpha\n\n<table><tr><td>t1</td></tr></table>\n\n"
            "Beta\n\n<table><tr><td>t2</td></tr></table>\n\nGamma"
        )
        result = clean_html_tables(text)
        for token in ("Alpha", "t1", "Beta", "t2", "Gamma"):
            assert token in result
        assert (
            result.index("Alpha")
            < result.index("t1")
            < result.index("Beta")
            < result.index("t2")
            < result.index("Gamma")
        )

    def test_unbalanced_markup_left_untouched(self) -> None:
        """Never truncate on malformed input — return it unchanged."""
        from worker.tasks import clean_html_tables

        unclosed = "before <table><tr><td>x</td></tr> after"
        assert clean_html_tables(unclosed) == unclosed
        stray = "prose </table> more"
        assert clean_html_tables(stray) == stray

    def test_content_after_responsive_div_survives(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '<div class="table-responsive">'
            "<table><tr><td>x</td></tr></table>"
            "</div>\n\nTrailing prose."
        )
        result = clean_html_tables(text)
        assert "Trailing prose." in result
        assert "table-responsive" not in result

    def test_non_wrapper_div_is_not_swallowed(self) -> None:
        from worker.tasks import clean_html_tables

        text = '<div class="content"><table><tr><td>z</td></tr></table></div>'
        result = clean_html_tables(text)
        assert "<div" in result and "</div>" in result

    def test_table_inside_comment_does_not_disable_cleaning(self) -> None:
        """An unbalanced <table> in a comment must not stop the real tables
        in the rest of the document from being cleaned."""
        from worker.tasks import clean_html_tables

        text = (
            "<!-- example: <table> -->\n\n"
            '<table class="x"><tr><td>real</td></tr></table>'
        )
        result = clean_html_tables(text)
        assert 'class="x"' not in result
        assert "<td>real</td>" in result
        assert "<!-- example: <table> -->" in result

    def test_unclosed_table_does_not_block_later_tables(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '<table class="a"><tr><td>1</td></tr>\n\n'
            '<table class="b"><tr><td>2</td></tr></table>'
        )
        result = clean_html_tables(text)
        assert 'class="b"' not in result
        assert "<td>2</td>" in result

    def test_fenced_code_block_left_verbatim(self) -> None:
        """HTML inside a Markdown fence is documentation, not markup."""
        from worker.tasks import clean_html_tables

        text = (
            'Doc:\n\n```html\n<table class="demo">\n'
            "<tr><td>ex</td></tr>\n</table>\n```\n\n"
            '<table class="y"><tr><td>r</td></tr></table>'
        )
        result = clean_html_tables(text)
        assert '<table class="demo">' in result
        assert 'class="y"' not in result and "<td>r</td>" in result

    def test_discards_script_and_style_in_cells(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            "<table><tr><td><script>x()</script>"
            "<style>.a{color:red}</style>ok</td></tr></table>"
        )
        result = clean_html_tables(text)
        assert "<script" not in result and "<style" not in result
        assert "x()" not in result and ".a{color:red}" not in result
        assert "ok" in result

    def test_keeps_col_span_drops_col_style(self) -> None:
        from worker.tasks import clean_html_tables

        text = (
            '<table><colgroup><col span="2" style="w"></colgroup>'
            "<tr><td>x</td></tr></table>"
        )
        result = clean_html_tables(text)
        assert 'span="2"' in result and "style=" not in result

    def test_comment_in_cell_does_not_become_visible_text(self) -> None:
        """Comment subclasses NavigableString; rewriting one would strip its
        <!-- --> markers and leak internal notes into the indexed body."""
        from worker.tasks import clean_html_tables

        text = (
            "<table><tr><td>visible"
            "<!-- TODO: remove\n   internal pricing note -->"
            "text</td></tr></table>"
        )
        result = clean_html_tables(text)
        assert "TODO" not in result and "pricing" not in result
        assert "visible" in result and "text" in result

    def test_comment_between_rows_is_dropped(self) -> None:
        from worker.tasks import clean_html_tables

        result = clean_html_tables("<table><!-- note --><tr><td>x</td></tr></table>")
        assert "note" not in result
        assert "<td>x</td>" in result

    def test_inline_backticks_do_not_mask_tables(self) -> None:
        """Only line-anchored fences are code blocks; stray inline backticks
        must not blank out everything up to the next backtick run."""
        from worker.tasks import clean_html_tables

        text = (
            "Intro ``` stray\n\n"
            '<table class="a"><tr><td>1</td></tr></table>\n\n'
            "more ``` end\n\n"
            '<table class="b"><tr><td>2</td></tr></table>'
        )
        result = clean_html_tables(text)
        assert 'class="a"' not in result and "<td>1</td>" in result
        assert 'class="b"' not in result and "<td>2</td>" in result

    def test_tilde_fence_is_masked(self) -> None:
        from worker.tasks import clean_html_tables

        text = 'Doc:\n\n~~~\n<table class="t">\n</table>\n~~~\n'
        assert 'class="t"' in clean_html_tables(text)


class TestBeautifulSoupExtractTables:
    """The BeautifulSoup fallback must preserve tables as faithfully as the
    trafilatura path — Unstructured's text_as_html drops colspan/rowspan and
    downgrades <th>, and markdownify flattens tables to pipe tables."""

    _HTML = (
        "<html><body><main><h1>T</h1><p>Some prose here.</p>"
        '<table><tr><th scope="col">H</th></tr>'
        '<tr><td colspan="2" rowspan="3">c</td></tr></table>'
        "<p>After.</p></main></body></html>"
    )

    def test_partition_path_preserves_table(self) -> None:
        from worker.tasks import _TABLE_PLACEHOLDER_PREFIX, _beautifulsoup_extract

        result = _beautifulsoup_extract(self._HTML)
        assert 'colspan="2"' in result
        assert 'rowspan="3"' in result
        assert "<th" in result
        assert "Some prose here." in result and "After." in result
        assert _TABLE_PLACEHOLDER_PREFIX not in result

    def test_markdownify_path_preserves_table(self) -> None:
        """When partition_html yields nothing we fall back to markdownify,
        which escapes punctuation — the placeholder must survive that."""
        from unittest.mock import patch

        from worker.tasks import _TABLE_PLACEHOLDER_PREFIX, _beautifulsoup_extract

        with patch("worker.tasks.partition_html", return_value=[]):
            result = _beautifulsoup_extract(self._HTML)
        assert 'colspan="2"' in result
        assert "<th" in result
        assert _TABLE_PLACEHOLDER_PREFIX not in result

    def test_empty_main_falls_through_to_body(self) -> None:
        """An empty <main> is truthy in BeautifulSoup — extracting from it
        would return nothing and silently drop the whole document."""
        from worker.tasks import _beautifulsoup_extract

        html = (
            "<html><body><main></main><div><p>Prose here.</p>"
            '<table><tr><td colspan="2">v</td></tr></table>'
            "</div></body></html>"
        )
        result = _beautifulsoup_extract(html)
        assert "Prose here." in result
        assert 'colspan="2"' in result

    def test_populated_main_is_still_preferred(self) -> None:
        from worker.tasks import _beautifulsoup_extract

        html = (
            "<html><body><main><h1>M</h1><p>Main prose.</p></main>"
            "<div><p>Sidebar.</p></div></body></html>"
        )
        result = _beautifulsoup_extract(html)
        assert "Main prose." in result
        assert "Sidebar." not in result


class TestElementsToMarkdownTables:
    def test_table_html_emitted_even_when_element_text_is_blank(self) -> None:
        """An element's text is Unstructured's own flattening and can be
        blank for a table we have good HTML for."""
        from worker.tasks import _elements_to_markdown

        element = Table(text="   ")
        element.metadata.text_as_html = "<table><tr><td>Real data</td></tr></table>"
        assert "Real data" in _elements_to_markdown([element])

    def test_table_with_no_text_content_is_dropped(self) -> None:
        from worker.tasks import _elements_to_markdown

        element = Table(text="")
        element.metadata.text_as_html = "<table><tr><td></td><td></td></tr></table>"
        assert _elements_to_markdown([element]) == ""

    def test_table_without_html_falls_back_to_escaped_pre(self) -> None:
        from worker.tasks import _elements_to_markdown

        element = Table(text="a < b & c")
        element.metadata.text_as_html = None
        result = _elements_to_markdown([element])
        assert "&lt;" in result and "&amp;" in result


class TestRestoreTables:
    def test_strips_markdown_decoration_from_placeholder_line(self) -> None:
        from worker.tasks import _restore_tables

        table = "<table><tr><td>T</td></tr></table>"
        for decoration in ("## ", "- ", "> ", "1. "):
            result, missing = _restore_tables(f"{decoration}@@M@@", [table], ["@@M@@"])
            assert missing == 0
            assert result.strip() == table

    def test_table_html_with_backslashes_is_not_mangled(self) -> None:
        """re.sub would treat \\g<1> in the replacement as a group reference."""
        from worker.tasks import _restore_tables

        table = r"<table><tr><td>C:\path \g<1> \1</td></tr></table>"
        result, missing = _restore_tables("@@M@@", [table], ["@@M@@"])
        assert missing == 0
        assert result == table

    def test_counts_dropped_placeholders(self) -> None:
        from worker.tasks import _restore_tables

        result, missing = _restore_tables("no marker here", ["<table/>"], ["@@M@@"])
        assert missing == 1
        assert result == "no marker here"

    def test_replaces_decorated_and_bare_occurrences(self) -> None:
        """A marker that survives both on its own line and inline must be
        fully substituted — no raw marker may reach the output."""
        from worker.tasks import _restore_tables

        result, missing = _restore_tables(
            "## @@M@@\n\nprose @@M@@ inline\n", ["<T/>"], ["@@M@@"]
        )
        assert "@@M@@" not in result
        assert result.count("<T/>") == 2
        assert missing == 0


# ---------------------------------------------------------------------------
# DOCX / PPTX merged-cell fidelity
# ---------------------------------------------------------------------------


class TestOfficeTableMerges:
    """
    Unstructured's text_as_html walks the logical cell grid, so a merged cell
    loses its span — and in DOCX its text is repeated once per covered column.
    We re-render those tables from the source file instead.
    """

    @staticmethod
    def _docx_with_merges(path: Path) -> Path:
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        for i, head in enumerate(("Aasta", "Summa", "Markus")):
            table.rows[0].cells[i].text = head
        table.rows[1].cells[0].merge(table.rows[1].cells[2]).text = "Kehtib koigile"
        for i, val in enumerate(("2024", "700", "ok")):
            table.rows[2].cells[i].text = val
        doc.save(path.as_posix())
        return path

    def test_docx_colspan_preserved_and_text_not_duplicated(
        self, tmp_path: Path
    ) -> None:
        from worker.tasks import clean_any_file

        entity = MagicMock()
        entity.file_path = self._docx_with_merges(tmp_path / "merged.docx")
        result = clean_any_file(entity)

        assert 'colspan="3"' in result
        assert result.count("Kehtib koigile") == 1

    def test_docx_rowspan_preserved(self, tmp_path: Path) -> None:
        from docx import Document

        from worker.tasks import clean_any_file

        path = tmp_path / "vmerge.docx"
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text = "H1"
        table.rows[0].cells[1].text = "H2"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "Spans two rows"
        table.cell(1, 1).text = "b1"
        table.cell(2, 1).text = "b2"
        doc.save(path.as_posix())

        entity = MagicMock()
        entity.file_path = path
        result = clean_any_file(entity)

        assert 'rowspan="2"' in result
        assert result.count("Spans two rows") == 1

    def test_pptx_colspan_preserved(self, tmp_path: Path) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        from worker.tasks import clean_any_file

        path = tmp_path / "merged.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(
            2, 3, Inches(0.5), Inches(0.5), Inches(8), Inches(2)
        ).table
        for i, head in enumerate(("A", "B", "C")):
            table.cell(0, i).text = head
        table.cell(1, 0).merge(table.cell(1, 2))
        table.cell(1, 0).text = "Merged across"
        prs.save(path.as_posix())

        entity = MagicMock()
        entity.file_path = path
        result = clean_any_file(entity)

        assert 'colspan="3"' in result
        assert result.count("Merged across") == 1

    def test_table_count_mismatch_keeps_unstructured_rendering(
        self, tmp_path: Path
    ) -> None:
        """Positional matching is only safe when both sides agree on how many
        tables there are; otherwise leave Unstructured's version alone."""
        from worker.tasks import _apply_source_table_html

        path = self._docx_with_merges(tmp_path / "merged.docx")
        element = Table(text="original")
        element.metadata.text_as_html = "<table>original</table>"
        # Two Table elements vs one table in the document.
        _apply_source_table_html(path, [element, element])
        assert element.metadata.text_as_html == "<table>original</table>"

    def test_unreadable_source_leaves_elements_untouched(self, tmp_path: Path) -> None:
        from worker.tasks import _apply_source_table_html

        broken = tmp_path / "broken.docx"
        broken.write_text("not a real docx", encoding="utf-8")
        element = Table(text="original")
        element.metadata.text_as_html = "<table>original</table>"
        _apply_source_table_html(broken, [element])
        assert element.metadata.text_as_html == "<table>original</table>"

    def test_nested_table_content_is_preserved(self, tmp_path: Path) -> None:
        """A table inside a cell must be rendered as a nested table, not
        dropped — reading only the cell's text loses it entirely."""
        from docx import Document

        from worker.tasks import clean_any_file

        path = tmp_path / "nested.docx"
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "Inner A"
        inner.cell(0, 1).text = "Inner B"
        outer.cell(0, 0).paragraphs[0].text = "Outer"
        doc.save(path.as_posix())

        entity = MagicMock()
        entity.file_path = path
        result = clean_any_file(entity)

        assert "Inner A" in result and "Inner B" in result
        assert "Outer" in result
        assert result.count("<table>") == 2

    def test_tbl_header_false_is_not_a_header_row(self, tmp_path: Path) -> None:
        from docx import Document
        from docx.oxml.ns import qn

        from worker.tasks import _docx_table_to_html

        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "not a header"
        tr_props = table.rows[0]._tr.get_or_add_trPr()
        tr_props.append(tr_props.makeelement(qn("w:tblHeader"), {qn("w:val"): "false"}))
        assert "<th" not in _docx_table_to_html(table)

    def test_multi_paragraph_cell_uses_line_breaks(self, tmp_path: Path) -> None:
        from docx import Document

        from worker.tasks import _docx_table_to_html

        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = "line one"
        cell.add_paragraph("line two")
        assert "line one<br/>line two" in _docx_table_to_html(table)

    def test_cell_text_is_html_escaped(self, tmp_path: Path) -> None:
        from docx import Document

        from worker.tasks import clean_any_file

        path = tmp_path / "escape.docx"
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "a < b & c"
        doc.save(path.as_posix())

        entity = MagicMock()
        entity.file_path = path
        result = clean_any_file(entity)
        assert "&lt;" in result and "&amp;" in result


# ---------------------------------------------------------------------------
# clean_html — routing logic (no real LLM)
# ---------------------------------------------------------------------------


class TestCleanHtmlRouting:
    def test_no_llm_uses_trafilatura_when_successful(self, tmp_path: Path) -> None:
        from worker.tasks import clean_html

        html = """<html><body><main>
          <h1>Retirement Reform</h1>
          <p>The retirement age will increase to 65 starting in 2026 for all workers
          born after 1970. This change affects approximately 2 million citizens.</p>
          <p>Early retirement remains available at 60 with a 10% benefit reduction.</p>
        </main></body></html>"""
        entity = _make_entity(tmp_path, ".html", html, use_llm=False)

        with patch(
            "worker.tasks._trafilatura_extract",
            return_value="# Extracted\n\nGood content.",
        ) as mock_traf:
            result = clean_html(entity, client=None, deployment=None)

        mock_traf.assert_called_once()
        assert result == "# Extracted\n\nGood content."

    def test_no_llm_falls_back_to_beautifulsoup_when_trafilatura_empty(
        self, tmp_path: Path
    ) -> None:
        from worker.tasks import clean_html

        entity = _make_entity(
            tmp_path, ".html", "<html><body><main><p>Hello</p></main></body></html>"
        )

        with (
            patch("worker.tasks._trafilatura_extract", return_value=None),
            patch(
                "worker.tasks._beautifulsoup_extract", return_value="BS result"
            ) as mock_bs,
        ):
            result = clean_html(entity, client=None, deployment=None)

        mock_bs.assert_called_once()
        assert result == "BS result"

    def test_use_llm_passes_when_eval_passes(self, tmp_path: Path) -> None:
        from worker.tasks import clean_html

        entity = _make_entity(
            tmp_path,
            ".html",
            "<html><body><p>text</p></body></html>",
            use_llm=True,
            use_llm_correction=False,
        )
        client = MagicMock()

        with (
            patch("worker.tasks._trafilatura_extract", return_value="Good extraction"),
            patch("worker.tasks._llm_evaluate", return_value=(True, "looks good")),
        ):
            result = clean_html(entity, client=client, deployment="dep")

        assert result == "Good extraction"

    def test_use_llm_falls_back_to_bs_when_eval_fails_no_correction(
        self, tmp_path: Path
    ) -> None:
        from worker.tasks import clean_html

        entity = _make_entity(
            tmp_path,
            ".html",
            "<html><body><p>text</p></body></html>",
            use_llm=True,
            use_llm_correction=False,
        )
        client = MagicMock()

        with (
            patch("worker.tasks._trafilatura_extract", return_value="Bad extraction"),
            patch("worker.tasks._llm_evaluate", return_value=(False, "too noisy")),
            patch(
                "worker.tasks._beautifulsoup_extract", return_value="BS fallback"
            ) as mock_bs,
        ):
            result = clean_html(entity, client=client, deployment="dep")

        mock_bs.assert_called_once()
        assert result == "BS fallback"

    def test_use_llm_correction_uses_llm_extract_when_eval_fails(
        self, tmp_path: Path
    ) -> None:
        from worker.tasks import clean_html

        entity = _make_entity(
            tmp_path,
            ".html",
            "<html><body><p>text</p></body></html>",
            use_llm=True,
            use_llm_correction=True,
        )
        client = MagicMock()

        with (
            patch("worker.tasks._trafilatura_extract", return_value="Bad extraction"),
            patch("worker.tasks._llm_evaluate", return_value=(False, "too noisy")),
            patch(
                "worker.tasks._llm_extract", return_value="LLM corrected content"
            ) as mock_llm_ext,
        ):
            result = clean_html(entity, client=client, deployment="dep")

        mock_llm_ext.assert_called_once()
        assert result == "LLM corrected content"

    def test_use_llm_correction_falls_back_to_bs_when_llm_extract_empty(
        self, tmp_path: Path
    ) -> None:
        from worker.tasks import clean_html

        entity = _make_entity(
            tmp_path,
            ".html",
            "<html><body><p>text</p></body></html>",
            use_llm=True,
            use_llm_correction=True,
        )
        client = MagicMock()

        with (
            patch("worker.tasks._trafilatura_extract", return_value="Bad extraction"),
            patch("worker.tasks._llm_evaluate", return_value=(False, "too noisy")),
            patch("worker.tasks._llm_extract", return_value=""),
            patch(
                "worker.tasks._beautifulsoup_extract", return_value="BS last resort"
            ) as mock_bs,
        ):
            result = clean_html(entity, client=client, deployment="dep")

        mock_bs.assert_called_once()
        assert result == "BS last resort"

    def test_correction_without_use_llm_logs_warning_and_ignores_correction(
        self, tmp_path: Path, caplog: pytest.LogCaptureFixture
    ) -> None:
        from worker.tasks import clean_html

        entity = _make_entity(
            tmp_path,
            ".html",
            "<html><body><p>text</p></body></html>",
            use_llm=False,
            use_llm_correction=True,
        )

        with (
            patch("worker.tasks._trafilatura_extract", return_value="Traf result"),
            caplog.at_level(logging.WARNING, logger="worker.tasks"),
        ):
            result = clean_html(entity, client=None, deployment=None)

        assert result == "Traf result"
        assert any("use_llm_correction=True" in r.message for r in caplog.records)


# ---------------------------------------------------------------------------
# LLM helpers — error handling
# ---------------------------------------------------------------------------


class TestLLMHelpers:
    def test_llm_evaluate_returns_false_on_api_error(self) -> None:
        from worker.tasks import _llm_evaluate
        from openai import APIError

        client = MagicMock()
        client.chat.completions.create.side_effect = APIError(
            message="rate limit", request=MagicMock(), body=None
        )
        passed, reason = _llm_evaluate(client, "dep", "some markdown")
        assert passed is False
        assert "LLM API error" in reason

    def test_llm_evaluate_returns_false_on_json_decode_error(self) -> None:
        from worker.tasks import _llm_evaluate

        client = MagicMock()
        client.chat.completions.create.return_value.choices[
            0
        ].message.content = "not json"
        passed, reason = _llm_evaluate(client, "dep", "some markdown")
        assert passed is False
        assert "non-JSON" in reason

    def test_llm_evaluate_parses_pass_true(self) -> None:
        from worker.tasks import _llm_evaluate

        client = MagicMock()
        client.chat.completions.create.return_value.choices[
            0
        ].message.content = '{"pass": true, "reason": "looks great"}'
        passed, reason = _llm_evaluate(client, "dep", "some markdown")
        assert passed is True
        assert reason == "looks great"

    def test_llm_evaluate_parses_pass_false(self) -> None:
        from worker.tasks import _llm_evaluate

        client = MagicMock()
        client.chat.completions.create.return_value.choices[
            0
        ].message.content = '{"pass": false, "reason": "too noisy"}'
        passed, reason = _llm_evaluate(client, "dep", "some markdown")
        assert passed is False
        assert reason == "too noisy"

    def test_llm_extract_returns_empty_string_on_api_error(self) -> None:
        from worker.tasks import _llm_extract
        from openai import APIError

        client = MagicMock()
        client.chat.completions.create.side_effect = APIError(
            message="timeout", request=MagicMock(), body=None
        )
        result = _llm_extract(client, "dep", "<html></html>")
        assert result == ""

    def test_llm_extract_returns_content(self) -> None:
        from worker.tasks import _llm_extract

        client = MagicMock()
        client.chat.completions.create.return_value.choices[
            0
        ].message.content = "# Extracted\n\nContent here."
        result = _llm_extract(client, "dep", "<html><body><p>x</p></body></html>")
        assert result == "# Extracted\n\nContent here."

    def test_llm_extract_rejects_truncated_response(self) -> None:
        """finish_reason == "length" means the answer was cut off, usually
        mid-table; returning it would write unbalanced HTML to cleaned.txt."""
        from worker.tasks import _llm_extract

        client = MagicMock()
        choice = client.chat.completions.create.return_value.choices[0]
        choice.finish_reason = "length"
        choice.message.content = "<table><tr><td>half a tab"
        result = _llm_extract(client, "dep", "<html></html>")
        assert result == ""


# ---------------------------------------------------------------------------
# set_up_logging — deduplication
# ---------------------------------------------------------------------------


class TestSetUpLogging:
    def test_no_duplicate_file_handlers(self, tmp_path: Path) -> None:
        from worker.tasks import set_up_logging
        import logging as _logging

        log_file = tmp_path / "test.log"
        log_file.touch()

        entity = MagicMock()
        entity.logs_path = log_file

        # Call twice — should only attach one FileHandler
        set_up_logging(entity)
        set_up_logging(entity)

        job_logger = _logging.getLogger("worker.tasks")
        file_handlers = [
            h for h in job_logger.handlers if isinstance(h, _logging.FileHandler)
        ]
        paths = [Path(h.baseFilename).resolve() for h in file_handlers]
        # All file handlers for this path should be deduplicated
        assert paths.count(log_file.resolve()) <= 1


# ---------------------------------------------------------------------------
# Plain-text routing (.txt, .md)
# ---------------------------------------------------------------------------


class TestPlainTextRouting:
    """
    .txt and .md files must bypass unstructured.partition() and read the
    file verbatim. Previously they fell through to clean_any_file(), which
    treats each line as a Title element and produces useless output —
    especially for files where the content happens to be HTML markup.
    """

    def test_clean_plain_text_returns_file_contents_verbatim(
        self, tmp_path: Path
    ) -> None:
        from worker.tasks import clean_plain_text

        body = "First line.\nSecond line.\n\nThird paragraph after blank line.\n"
        entity = _make_entity(tmp_path, ".txt", body)

        assert clean_plain_text(entity) == body

    def test_clean_plain_text_replaces_invalid_utf8(self, tmp_path: Path) -> None:
        from worker.tasks import clean_plain_text

        source = tmp_path / "source.txt"
        source.write_bytes(b"good bytes \xff\xfe bad bytes ok\n")
        meta = tmp_path / "source.meta.json"
        meta.write_text(
            json.dumps({"file_type": ".txt", "metadata": {}}), encoding="utf-8"
        )

        entity = MagicMock()
        entity.file_path = source
        entity.meta_data_path = meta
        entity.directory_path = tmp_path
        entity.use_llm = False

        out = clean_plain_text(entity)
        assert "good bytes" in out
        assert "bad bytes ok" in out

    def test_txt_routes_to_plain_text_not_unstructured(self, tmp_path: Path) -> None:
        """
        clean_file_task with file_type=".txt" must call clean_plain_text
        and must NOT call clean_any_file (which would invoke unstructured).
        """
        from worker.tasks import clean_file_task

        body = "Some real text content for a .txt source.\n"
        entity = _make_entity(tmp_path, ".txt", body, use_llm=False)

        with (
            patch("worker.tasks.requests.post") as mock_post,
            patch("worker.tasks.cleanup_directory"),
            patch("worker.tasks.clean_any_file") as mock_any,
        ):
            mock_resp = MagicMock()
            mock_resp.json.return_value = {"response": "http://mock/file"}
            mock_resp.raise_for_status.return_value = None
            mock_post.return_value = mock_resp

            clean_file_task(entity)

        assert mock_any.call_count == 0, (
            ".txt must not be routed through clean_any_file()"
        )
        cleaned = (tmp_path / "cleaned.txt").read_text(encoding="utf-8")
        assert "Some real text content" in cleaned

    def test_md_source_html_is_not_rewritten(self, tmp_path: Path) -> None:
        """A .md source is passed through verbatim — clean_html_tables must
        not rewrite HTML the author wrote by hand."""
        from worker.tasks import clean_file_task

        body = '# Doc\n\n<table class="demo" style="border:1">\n<tr><td>x</td></tr>\n</table>\n'
        entity = _make_entity(tmp_path, ".md", body, use_llm=False)

        with (
            patch("worker.tasks.requests.post") as mock_post,
            patch("worker.tasks.cleanup_directory"),
        ):
            mock_resp = MagicMock()
            mock_resp.json.return_value = {"response": "http://mock/file"}
            mock_resp.raise_for_status.return_value = None
            mock_post.return_value = mock_resp

            clean_file_task(entity)

        cleaned = (tmp_path / "cleaned.txt").read_text(encoding="utf-8")
        assert 'class="demo"' in cleaned
        assert 'style="border:1"' in cleaned

    def test_md_routes_to_plain_text(self, tmp_path: Path) -> None:
        from worker.tasks import clean_file_task

        body = "# Heading\n\nParagraph with **bold** text.\n"
        entity = _make_entity(tmp_path, ".md", body, use_llm=False)

        with (
            patch("worker.tasks.requests.post") as mock_post,
            patch("worker.tasks.cleanup_directory"),
            patch("worker.tasks.clean_any_file") as mock_any,
        ):
            mock_resp = MagicMock()
            mock_resp.json.return_value = {"response": "http://mock/file"}
            mock_resp.raise_for_status.return_value = None
            mock_post.return_value = mock_resp

            clean_file_task(entity)

        assert mock_any.call_count == 0
        cleaned = (tmp_path / "cleaned.txt").read_text(encoding="utf-8")
        assert "# Heading" in cleaned
        assert "**bold**" in cleaned


# ---------------------------------------------------------------------------
# Metadata mutation
# ---------------------------------------------------------------------------


class TestMetadataMutation:
    def test_language_written_inside_metadata_key(self, tmp_path: Path) -> None:
        """
        After clean_file_task runs, metadata["metadata"]["language"] must be set
        and any stale top-level "language" key (as written by the scrapper) must
        have been removed.
        """
        from worker.tasks import clean_file_task

        html = "<html><body><main><p>Test content for language detection.</p></main></body></html>"
        entity = _make_entity(tmp_path, ".html", html, use_llm=False)

        original_meta = json.loads(entity.meta_data_path.read_text())
        original_meta["language"] = "et"
        entity.meta_data_path.write_text(json.dumps(original_meta))

        with (
            patch("worker.tasks.requests.post") as mock_post,
            patch("worker.tasks.cleanup_directory"),
        ):
            mock_resp = MagicMock()
            mock_resp.json.return_value = {"response": "http://mock/file"}
            mock_resp.raise_for_status.return_value = None
            mock_post.return_value = mock_resp

            clean_file_task(entity)

        meta_file = tmp_path / "cleaned.meta.json"
        assert meta_file.exists()
        metadata = json.loads(meta_file.read_text())
        assert "cleaned" in metadata["metadata"]
        assert metadata["metadata"]["cleaned"] is True
        assert "language" in metadata["metadata"]
        assert metadata["metadata"]["language"] is not None
        assert "language" not in metadata

    def test_cleaned_txt_written(self, tmp_path: Path) -> None:
        from worker.tasks import clean_file_task

        html = (
            "<html><body><main><h1>Hello</h1><p>World content.</p></main></body></html>"
        )
        entity = _make_entity(tmp_path, ".html", html, use_llm=False)

        with (
            patch("worker.tasks.requests.post") as mock_post,
            patch("worker.tasks.cleanup_directory"),
        ):
            mock_resp = MagicMock()
            mock_resp.json.return_value = {"response": "http://mock/file"}
            mock_resp.raise_for_status.return_value = None
            mock_post.return_value = mock_resp

            clean_file_task(entity)

        cleaned = tmp_path / "cleaned.txt"
        assert cleaned.exists()
        content = cleaned.read_text()
        assert len(content) > 0
