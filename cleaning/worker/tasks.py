import base64
import datetime
import ipaddress
import json
import logging
import mimetypes
import re
import socket
import textwrap
from html import escape as html_escape
from pathlib import Path
from urllib.parse import urljoin, urlparse
from uuid import uuid4

import pymupdf4llm
import requests
import trafilatura
from bs4 import BeautifulSoup
from bs4.element import NavigableString, PreformattedString, Tag
from langdetect import detect, LangDetectException
from markdownify import markdownify
from openai import AzureOpenAI, APIError

from unstructured.documents.elements import Title, ListItem, Table, CodeSnippet
from unstructured.partition.auto import partition
from unstructured.partition.html import partition_html

from api.config import settings, get_vault_secrets, VaultSecrets
from api.models import EntityToClean, SourceCleaningTask
from worker.utils import catch_error, cleanup_directory, send_error

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# elements_to_markdown helper
# ---------------------------------------------------------------------------
# unstructured.staging.base.elements_to_markdown was removed in 0.18.x.
# We implement the same logic directly using the element types we care about.


# Strips tags to test whether a fragment carries any actual text.
_TAG_RE = re.compile(r"<[^>]+>")


def _elements_to_markdown(elements: list) -> str:
    """
    Convert a list of Unstructured elements to a Markdown string.

    Mapping:
      Title       -> ## heading
      ListItem    -> - bullet
      Table       -> raw HTML from Unstructured's text_as_html metadata.
                     Markdown pipe tables cannot represent colspan/rowspan
                     or multi-line cells; keeping HTML preserves the source.
      CodeSnippet -> fenced code block
      Everything else -> plain paragraph
    """
    lines: list[str] = []
    for el in elements:
        text = str(el).strip()
        if isinstance(el, Table):
            # Checked before the empty-text guard: an element's text is
            # Unstructured's own flattening, which can come back blank for a
            # table we have perfectly good HTML for.
            table_html = getattr(el.metadata, "text_as_html", None)
            if table_html and _TAG_RE.sub("", table_html).strip():
                lines.append(table_html)
            elif text:
                lines.append(f"<pre>{html_escape(text)}</pre>")
            continue
        if not text:
            continue
        if isinstance(el, Title):
            lines.append(f"## {text}")
        elif isinstance(el, ListItem):
            lines.append(f"- {text}")
        elif isinstance(el, CodeSnippet):
            lines.append(f"```\n{text}\n```")
        else:
            lines.append(text)
    return "\n\n".join(lines)


# ---------------------------------------------------------------------------
# Azure OpenAI client -- created per task using fresh Vault secrets
# ---------------------------------------------------------------------------


def _make_openai_client(secrets: VaultSecrets) -> AzureOpenAI:
    return AzureOpenAI(
        api_key=secrets.azure_openai_api_key.get_secret_value(),
        api_version=secrets.azure_openai_api_version,
        azure_endpoint=secrets.azure_openai_endpoint,
    )


# ---------------------------------------------------------------------------
# Text normalisation
# ---------------------------------------------------------------------------


def normalize_newlines(text: str) -> str:
    """Collapse 3+ consecutive newlines down to 2."""
    return re.sub(r"\n{3,}", "\n\n", text)


# ---------------------------------------------------------------------------
# HTML table cleanup — strip presentational cruft while keeping structure
# ---------------------------------------------------------------------------

# Attributes preserved on table elements — everything else is removed.
_KEEP_TABLE_ATTRS: dict[str, set[str]] = {
    "table": {"summary"},
    "th": {"colspan", "rowspan", "scope"},
    "td": {"colspan", "rowspan"},
    "a": {"href", "title"},
    "img": {"src", "alt", "title"},
    "col": {"span"},
    "colgroup": {"span"},
}
# Tags unwrapped inside tables (pure styling wrappers) — content is kept.
_DROP_TAGS_IN_TABLE: set[str] = {"font", "span"}
# Tags removed inside tables along with their contents — never document text.
_DISCARD_TAGS_IN_TABLE: set[str] = {"script", "style", "noscript"}


def _clean_table_element(el: object) -> None:
    """Recursively strip presentational attrs from a table subtree, drop
    script/style, unwrap span/font, downgrade headings inside cells to
    <strong>, flatten nested <strong>.
    """
    if not isinstance(el, Tag):
        return

    for child in list(el.children):
        # Comments, CDATA and processing instructions are not document text.
        # They must be removed rather than left to the serialiser, which would
        # otherwise rewrite them into visible cell content.
        if isinstance(child, PreformattedString):
            child.extract()
            continue
        if not isinstance(child, Tag):
            continue
        if child.name in _DISCARD_TAGS_IN_TABLE:
            child.decompose()
            continue
        if child.name in _DROP_TAGS_IN_TABLE:
            _clean_table_element(child)
            child.unwrap()
            continue
        _clean_table_element(child)

    keep = _KEEP_TABLE_ATTRS.get(el.name, set())
    for attr in list(el.attrs.keys()):
        if attr not in keep:
            del el.attrs[attr]

    if el.name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        el.name = "strong"

    if el.name == "strong":
        for inner in el.find_all("strong"):
            if isinstance(inner, Tag):
                inner.unwrap()


# Tags whose whitespace-only text children are pure indentation: HTML never
# renders whitespace that sits in a table outside a cell, so it can be dropped.
# Whitespace inside <td>/<th> (or inline elements) IS significant and is kept.
_TABLE_STRUCTURE_TAGS: set[str] = {
    "table",
    "thead",
    "tbody",
    "tfoot",
    "tr",
    "colgroup",
}

# Structural tags placed on their own line when serialising, for readability.
_TABLE_LINE_BREAK_RE = re.compile(r"(</?(?:thead|tbody|tfoot|tr)>)")

# Matches any <table ...> / </table> tag, used to find balanced table spans.
_TABLE_TAG_RE = re.compile(r"<(/?)table\b[^>]*>", re.IGNORECASE)

# Regions that look like markup but are not: HTML comments and fenced code
# blocks. They are blanked out (length-preserving, so offsets still line up
# with the original text) before scanning, so a <table> written inside a
# comment or a ```html example is neither rewritten nor allowed to unbalance
# the depth counter for the rest of the document.
_MASKED_REGION_RE = re.compile(
    # An HTML comment, or a fenced code block whose delimiters both start a
    # line. Anchoring to line starts keeps stray inline backticks from
    # masking everything up to the next backtick run.
    r"<!--.*?-->|^[ \t]*(```|~~~)[^\n]*$.*?^[ \t]*\1",
    re.DOTALL | re.MULTILINE,
)

# Malformed input safety valve: how many unclosed <table> tags to neutralise
# before giving up, so one broken table cannot disable cleaning document-wide.
_MAX_SPAN_RECOVERIES = 8

# A <div class="... table-responsive ..."> wrapper immediately before a table
# (and its matching </div> after) is Bootstrap scroll scaffolding, not content.
_RESPONSIVE_DIV_OPEN_RE = re.compile(
    r"<div\b[^>]*\bclass\s*=\s*(\"[^\"]*\"|\'[^\']*\'|[^\s>]+)[^>]*>\s*$",
    re.IGNORECASE,
)
_RESPONSIVE_DIV_CLOSE_RE = re.compile(r"\s*</div>", re.IGNORECASE)


def _mask_non_markup(text: str) -> str:
    """Blank out comments and fenced code blocks, preserving length and
    newlines so offsets into the result also index the original text."""
    return _MASKED_REGION_RE.sub(lambda m: re.sub(r"[^\n]", " ", m.group(0)), text)


def _scan_table_spans(
    masked: str,
) -> tuple[list[tuple[int, int]], tuple[int, int] | None]:
    """
    Scan for balanced <table> spans, tracking nesting depth.

    Returns (spans, unclosed) where *unclosed* is the span of the outermost
    <table> tag that is never closed, or None when the markup is balanced.
    """
    spans: list[tuple[int, int]] = []
    open_tags: list[tuple[int, int]] = []
    start = 0
    for match in _TABLE_TAG_RE.finditer(masked):
        if match.group(1):  # closing tag
            if not open_tags:
                continue  # stray </table>, nothing open
            open_tags.pop()
            if not open_tags:
                spans.append((start, match.end()))
        else:
            if not open_tags:
                start = match.start()
            open_tags.append(match.span())
    return spans, open_tags[0] if open_tags else None


def _find_table_spans(text: str) -> list[tuple[int, int]]:
    """
    Return (start, end) offsets of every top-level <table>...</table> block.

    Nesting is tracked by depth so an outer table's span covers any tables
    nested inside it — a non-greedy regex would stop at the first </table>
    and cut the outer table in half. Nested tables are not returned
    separately; they are cleaned as part of their parent's subtree.

    An unclosed <table> is neutralised and the scan retried, so malformed
    markup costs only that one table rather than every table after it.
    """
    masked = _mask_non_markup(text)
    spans: list[tuple[int, int]] = []
    for _ in range(_MAX_SPAN_RECOVERIES):
        spans, unclosed = _scan_table_spans(masked)
        if unclosed is None:
            return spans
        open_start, open_end = unclosed
        logger.warning(
            f"[tables] unclosed <table> at offset {open_start}; "
            "leaving it as-is and continuing"
        )
        masked = masked[:open_start] + " " * (open_end - open_start) + masked[open_end:]
    return spans


def _expand_over_responsive_div(text: str, start: int, end: int) -> tuple[int, int]:
    """Widen a table span to swallow a wrapping table-responsive <div>."""
    open_match = _RESPONSIVE_DIV_OPEN_RE.search(text, 0, start)
    if not open_match or "table-responsive" not in open_match.group(1).lower():
        return start, end
    close_match = _RESPONSIVE_DIV_CLOSE_RE.match(text, end)
    if not close_match:
        return start, end
    return open_match.start(), close_match.end()


def _serialise_table(table: Tag) -> str:
    """
    Render a cleaned table as HTML with structural tags on their own lines.

    Whitespace is normalised on the parse tree rather than on the serialised
    string: runs of whitespace inside each text node collapse to one space,
    and whitespace-only nodes are dropped only where they are indentation
    (directly inside table/thead/tbody/tfoot/tr). Collapsing on the serialised
    string instead would delete the single space in "</strong> <em>", silently
    welding adjacent words together.
    """
    for node in list(table.descendants):
        # Comment and friends subclass NavigableString; replacing one with a
        # plain NavigableString would strip its <!-- --> markers and leak the
        # comment into the document body.
        if isinstance(node, PreformattedString):
            node.extract()
            continue
        if not isinstance(node, NavigableString):
            continue
        parent = node.parent
        collapsed = re.sub(r"\s+", " ", str(node))
        if (
            not collapsed.strip()
            and parent is not None
            and parent.name in _TABLE_STRUCTURE_TAGS
        ):
            node.extract()
            continue
        if collapsed != str(node):
            node.replace_with(NavigableString(collapsed))

    return _TABLE_LINE_BREAK_RE.sub(r"\n\1", str(table)).strip()


def clean_html_tables(text: str) -> str:
    """
    Strip presentational cruft from every <table>...</table> block found
    inside a Markdown document, leaving colspan/rowspan/scope/href/title
    intact. Content outside <table> blocks is untouched.

    Also re-formats each table so <thead>, <tbody> and every <tr> start on
    their own line — same rendered output, easier to read in raw form.
    """
    spans = _find_table_spans(text)
    if not spans:
        return text

    out: list[str] = []
    cursor = 0
    for start, end in spans:
        start, end = _expand_over_responsive_div(text, start, end)
        if start < cursor:  # wrapper already consumed by the previous span
            continue
        fragment = BeautifulSoup(text[start:end], "lxml")
        table = fragment.find("table")
        out.append(text[cursor:start])
        if not isinstance(table, Tag):
            out.append(text[start:end])
        else:
            _clean_table_element(table)
            out.append(_serialise_table(table))
        cursor = end
    out.append(text[cursor:])
    return "".join(out)


# ---------------------------------------------------------------------------
# HTML extraction helpers
# ---------------------------------------------------------------------------


# Placeholder tokens are deliberately alphanumeric: Markdown writers escape
# punctuation (markdownify turns "@@CKB_TABLE_1@@" into "@@CKB\\_TABLE\\_1@@"),
# which would stop the marker matching and silently lose the table.
_TABLE_PLACEHOLDER_PREFIX = "CKBTABLEPLACEHOLDER"
_TABLE_PLACEHOLDER_SUFFIX = "ENDCKBTABLE"

# A placeholder sitting alone on a line may pick up Markdown decoration from
# the extractor (a list bullet, a heading marker, a blockquote arrow) because
# it looked like a short standalone paragraph. Strip that off when restoring
# so the table HTML does not end up glued to a "## " or "- ".
_MARKDOWN_LINE_PREFIX = r"(?:[ \t]*(?:#{1,6}[ \t]+|[-*+][ \t]+|\d+\.[ \t]+|>[ \t]*))*"


def _swap_tables_for_placeholders(soup: BeautifulSoup) -> tuple[list[str], list[str]]:
    """
    Replace every top-level <table> in *soup* with a paragraph placeholder.

    Returns (table_html, markers) as parallel lists. Nested tables are left
    alone — they travel inside their parent's captured HTML. Markers carry a
    per-call nonce so text that happens to look like a placeholder in the
    source page cannot collide with a real one.
    """
    nonce = uuid4().hex[:8]
    tables: list[str] = []
    markers: list[str] = []
    for table in soup.find_all("table"):
        if table.find_parent("table") is not None:
            continue
        marker = (
            f"{_TABLE_PLACEHOLDER_PREFIX}{nonce}X{len(tables)}"
            f"{_TABLE_PLACEHOLDER_SUFFIX}"
        )
        tables.append(str(table))
        markers.append(marker)
        marker_tag = soup.new_tag("p")
        marker_tag.string = marker
        table.replace_with(marker_tag)
    return tables, markers


def _restore_tables(
    text: str, tables: list[str], markers: list[str]
) -> tuple[str, int]:
    """
    Substitute captured table HTML back in for its placeholder.

    Returns (text, missing) where *missing* counts placeholders the extractor
    dropped — normally boilerplate tables pruned as out-of-scope, but worth
    logging so a broken round-trip is distinguishable from correct pruning.
    """
    missing = 0
    for marker, table_html in zip(markers, tables, strict=True):
        # Tolerate a Markdown writer having backslash-escaped characters of
        # the marker on its way through the extractor.
        tolerant = r"\\?".join(map(re.escape, marker))
        patterns = (
            re.compile(rf"^{_MARKDOWN_LINE_PREFIX}{tolerant}[ \t]*$", re.MULTILINE),
            re.compile(tolerant),
        )
        replaced = 0
        for pattern in patterns:
            # A function replacement avoids re.sub treating backslashes and
            # \g<..> in the table HTML as escape sequences.
            text, count = pattern.subn(lambda _, repl=table_html: repl, text)
            replaced += count
        if not replaced:
            missing += 1
    return text, missing


def _beautifulsoup_extract(html: str) -> str:
    """
    Multi-step fallback extractor that produces Markdown output.

    Steps:
      1. Strip noisy elements (header, footer, nav, script, style, aside, form).
      2. Swap tables out for placeholders so neither Unstructured's
         text_as_html (which drops colspan/rowspan and downgrades <th>) nor
         markdownify's pipe-table renderer can flatten them.
      3. Extract from <main>, else <body>, else the cleaned root (avoiding
         html/head wrapper noise), via partition_html with
         skip_headers_and_footers, falling back to markdownify when
         partition returns nothing. Move to the next candidate if a target
         yields no text.
      4. Put the original table HTML back.
    """
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["header", "footer", "nav", "script", "style", "aside", "form"]):
        tag.decompose()

    tables, markers = _swap_tables_for_placeholders(soup)

    # Try <main>, then <body>, then the cleaned root. An empty <main> is
    # truthy in BeautifulSoup, so a page whose real content sits beside an
    # empty <main> would otherwise extract to nothing at all; falling through
    # on an empty result keeps that content instead of dropping the document.
    extracted = ""
    for target in (soup.find("main"), soup.find("body"), soup):
        if target is None:
            continue
        partitioned = partition_html(
            text=str(target),
            languages=settings.languages,
            skip_headers_and_footers=True,
        )
        if partitioned:
            extracted = _elements_to_markdown(partitioned)
        else:
            extracted = markdownify(str(target), heading_style="ATX")
        if extracted.strip():
            break

    extracted, missing = _restore_tables(extracted, tables, markers)
    if missing:
        logger.info(
            f"[html] BeautifulSoup extraction dropped {missing} of {len(tables)} "
            "table(s) as out-of-scope content"
        )
    return extracted


def _trafilatura_extract(html: str, url: str | None = None) -> str | None:
    """
    Extract main content as Markdown, preserving tables as raw HTML.

    Two-pass strategy: BeautifulSoup replaces each <table> with a plain-text
    placeholder before trafilatura runs; the original table HTML is then
    substituted back into trafilatura's Markdown output. This avoids
    trafilatura's lossy Markdown table renderer (which cannot represent
    colspan/rowspan or multi-line cells) while keeping trafilatura's strong
    prose extraction.
    """
    soup = BeautifulSoup(html, "lxml")
    tables, markers = _swap_tables_for_placeholders(soup)

    result = trafilatura.extract(
        str(soup) if tables else html,
        url=url,
        output_format="markdown",
        include_comments=False,
        include_tables=not tables,
        favor_precision=True,
    )
    if not result:
        return None

    result, missing = _restore_tables(result, tables, markers)
    if missing:
        logger.info(
            f"[html] trafilatura dropped {missing} of {len(tables)} table(s) "
            "as out-of-scope content"
        )
    return result or None


# ---------------------------------------------------------------------------
# LLM helpers
# ---------------------------------------------------------------------------

_EVAL_SYSTEM = textwrap.dedent("""
    You are a quality-evaluation assistant for web-page text extraction.
    You will receive a Markdown-formatted extraction of a web page's main body.
    Evaluate it on the following criteria:
      1. Readability  - is the text coherent and human-readable?
      2. Completeness - does it appear to contain the full main content
         without obvious truncation or missing sections?
      3. Cleanliness  - is it free from navigation menus, cookie banners,
         footer boilerplate, and other noise?
      4. Structure    - are headings, lists, and paragraphs logically preserved?

    Tables may appear as raw HTML (<table>, <tr>, <td>, colspan, rowspan).
    This is expected and should not count as "noise" or against cleanliness.

    Reply with ONLY a JSON object in this exact shape (no markdown fences):
    {"pass": true, "reason": "<one-sentence explanation>"}
    or
    {"pass": false, "reason": "<one-sentence explanation>"}
""").strip()

_EXTRACT_SYSTEM = textwrap.dedent("""
    You are an expert web-page content extractor.
    You will receive raw HTML of a web page.
    Extract ONLY the main body content (article text, tables, documentation).
    Ignore navigation, sidebars, footers, cookie notices, and ads.
    Return clean Markdown for prose (headings, lists, bold/italic where
    appropriate). Do not add any commentary of your own.

    Tables:
    - Preserve every row and column verbatim; do not drop, merge or
      paraphrase cell contents.
    - If a table uses colspan or rowspan, or has nested <ul>/<li> inside
      cells, output the entire table as HTML (<table>, <thead>, <tbody>,
      <tr>, <th>, <td colspan="N">, <td rowspan="N">). Keep <ul>/<li>,
      <strong>, <sup>, <br> and inline <a href="..."> inside cells verbatim.
    - Only if a table is simple (no colspan/rowspan, single-line cells)
      may you use a Markdown pipe table.
    - Never wrap tables in code fences.
    - Never include style, class, align, cellpadding, cellspacing, border,
      width or bgcolor attributes on table elements — only colspan, rowspan
      and scope carry meaning.
""").strip()


def _llm_evaluate(
    client: AzureOpenAI, deployment: str, extracted_markdown: str
) -> tuple[bool, str]:
    """
    Ask the LLM to evaluate the quality of an extraction.
    Returns (passed, reason). On any API or parse error returns (False, reason)
    so the caller can fall back gracefully without crashing.
    """
    try:
        response = client.chat.completions.create(
            model=deployment,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": _EVAL_SYSTEM},
                {
                    "role": "user",
                    "content": f"# Extraction to evaluate\n\n{extracted_markdown}",
                },
            ],
        )
        raw = response.choices[0].message.content or "{}"
    except APIError as e:
        return False, f"LLM API error during evaluation: {e}"
    except Exception as e:
        return False, f"Unexpected error during LLM evaluation: {e}"

    try:
        result = json.loads(raw)
        return bool(result.get("pass", False)), result.get("reason", "no reason given")
    except json.JSONDecodeError:
        return False, f"LLM returned non-JSON: {raw[:120]}"


def _llm_extract(client: AzureOpenAI, deployment: str, html: str) -> str:
    """
    Ask the LLM to re-extract the main content from raw HTML.
    Returns the extracted Markdown, or empty string on any error so the
    caller can fall back gracefully without crashing.
    """
    try:
        response = client.chat.completions.create(
            model=deployment,
            messages=[
                {"role": "system", "content": _EXTRACT_SYSTEM},
                {"role": "user", "content": html},
            ],
        )
        choice = response.choices[0]
        # A response cut off at the output limit usually breaks mid-table,
        # leaving unbalanced HTML. Treat it as a failure so the caller falls
        # back rather than writing malformed markup to cleaned.txt.
        if choice.finish_reason == "length":
            logger.error(
                "LLM re-extraction hit the output token limit and was truncated"
            )
            return ""
        return (choice.message.content or "").strip()
    except APIError as e:
        logger.error(f"LLM API error during re-extraction: {e}")
        return ""
    except Exception as e:
        logger.error(f"Unexpected error during LLM re-extraction: {e}")
        return ""


# ---------------------------------------------------------------------------
# HTML cleaning
# ---------------------------------------------------------------------------


def clean_html(
    entity: EntityToClean, client: AzureOpenAI | None, deployment: str | None
) -> str:
    """
    Three modes controlled by entity.use_llm and entity.use_llm_correction:

    use_llm=False (default):
        trafilatura -> on empty/error fall back to BeautifulSoup.

    use_llm=True, use_llm_correction=False:
        trafilatura (or BeautifulSoup if trafilatura is empty) ->
        LLM evaluation -> if PASS return extraction, if FAIL fall back
        to BeautifulSoup.

    use_llm=True, use_llm_correction=True:
        trafilatura (or BeautifulSoup if trafilatura is empty) ->
        LLM evaluation -> if PASS return extraction, if FAIL send raw
        HTML to LLM for full re-extraction -> if LLM re-extraction is
        empty fall back to BeautifulSoup.

    Note: use_llm_correction=True has no effect when use_llm=False.
    """
    with entity.file_path.open("r", encoding="utf-8", errors="replace") as f:
        html = f.read()

    # Guard: warn if correction is requested without evaluation being enabled.
    if entity.use_llm_correction and not entity.use_llm:
        logger.warning(
            f"[html] use_llm_correction=True but use_llm=False for {entity.url}; "
            "correction will be ignored -- set use_llm=True to enable it."
        )

    # No LLM path
    if not entity.use_llm:
        extracted = _trafilatura_extract(html, url=entity.url)
        if extracted:
            logger.info(f"[html] trafilatura succeeded for {entity.url}")
            return extracted
        logger.warning(
            f"[html] trafilatura empty, falling back to BeautifulSoup for {entity.url}"
        )
        return _beautifulsoup_extract(html)

    # LLM paths — entity.use_llm is True here, so the caller must have provided
    # an Azure OpenAI client + deployment; assert to narrow the optional types.
    assert client is not None and deployment is not None, (
        "use_llm=True requires Azure OpenAI client and deployment"
    )

    extracted = _trafilatura_extract(html, url=entity.url)
    if not extracted:
        logger.warning(
            f"[html] trafilatura empty for {entity.url}; using BeautifulSoup before LLM eval"
        )
        extracted = _beautifulsoup_extract(html)

    logger.info(f"[html] running LLM evaluation for {entity.url}")
    passed, reason = _llm_evaluate(client, deployment, extracted)
    logger.info(
        f"[html] LLM evaluation {'PASSED' if passed else 'FAILED'} for {entity.url}: {reason}"
    )

    if passed:
        return extracted

    if not entity.use_llm_correction:
        logger.warning(
            f"[html] LLM eval failed, falling back to BeautifulSoup for {entity.url}"
        )
        return _beautifulsoup_extract(html)

    logger.info(f"[html] LLM correction: re-extracting from raw HTML for {entity.url}")
    corrected = _llm_extract(client, deployment, html)
    if corrected:
        return corrected

    logger.warning(
        f"[html] LLM re-extraction empty; falling back to BeautifulSoup for {entity.url}"
    )
    return _beautifulsoup_extract(html)


# ---------------------------------------------------------------------------
# Non-HTML cleaning
# ---------------------------------------------------------------------------


def clean_pdf(entity: EntityToClean) -> str:
    """
    Use pymupdf4llm for PDF-to-Markdown conversion.
    Handles multi-column layouts, tables, headings, and hyperlinks automatically.
    Images are intentionally skipped here — they are extracted separately by
    extract_images_from_pdf() and uploaded as individual files.
    """
    return pymupdf4llm.to_markdown(entity.file_path.as_posix(), show_progress=False)


def clean_any_file(entity: EntityToClean) -> str:
    """
    Use Unstructured for all non-HTML, non-PDF formats (DOCX, DOC, etc.).
    Output is rendered as Markdown via _elements_to_markdown so headings,
    lists, tables, and emphasis are all preserved.
    """
    partitioned = partition(
        filename=entity.file_path.as_posix(), languages=settings.languages
    )
    return _elements_to_markdown(partitioned)


# Sources that are already text: cleaned by passthrough, never rewritten.
_PLAIN_TEXT_TYPES: frozenset[str] = frozenset({".txt", ".md"})


def clean_plain_text(entity: EntityToClean) -> str:
    """
    Read a plain-text file (.txt / .md) as UTF-8 and return its contents.

    No extraction is performed — the file *is* the text. Whitespace
    normalisation is applied by the caller via normalize_newlines().
    Invalid bytes are replaced rather than raising so a single bad byte
    in a long file doesn't kill the whole job.
    """
    with entity.file_path.open("r", encoding="utf-8", errors="replace") as f:
        return f.read()


# ---------------------------------------------------------------------------
# Image extraction
# ---------------------------------------------------------------------------

# Shared HTTP session for HTML image downloads.
# A persistent session reuses the underlying TCP connection pool and ensures
# a consistent User-Agent header without repeating it on every call.
_HTML_IMAGE_SESSION = requests.Session()
_HTML_IMAGE_SESSION.headers.update(
    {
        "User-Agent": "Mozilla/5.0 (compatible; CKB-Cleaner/1.0)",
    }
)


def _is_private_url(url: str) -> bool:
    """
    Return True if *url* resolves to a non-globally-routable address.

    This is a basic SSRF guard: it prevents the worker from fetching resources
    on the internal network (loopback, RFC-1918 private ranges, link-local,
    etc.) when processing untrusted HTML <img src> attributes.

    Returns True (i.e. "block it") on any resolution failure so the default
    is to skip rather than to fetch when the host is ambiguous.
    """
    try:
        host = urlparse(url).hostname
        if not host:
            return True
        ip = ipaddress.ip_address(socket.gethostbyname(host))
        return not ip.is_global
    except Exception:
        return True


def _images_dir(entity: EntityToClean) -> Path:
    """Return (and create) the per-job images subdirectory."""
    d = entity.directory_path / "images"
    d.mkdir(exist_ok=True)
    return d


def extract_images_from_html(entity: EntityToClean) -> list[Path]:
    """
    Extract images referenced by <img src="..."> in the HTML file.

    - data-URI images are decoded and saved immediately.
    - http/https URLs are downloaded; relative URLs are resolved against
      entity.url before downloading.
    - Any individual download failure is logged and skipped so one bad
      image does not abort the whole job.
    """
    with entity.file_path.open("r", encoding="utf-8", errors="replace") as f:
        html = f.read()

    soup = BeautifulSoup(html, "lxml")
    images_dir = _images_dir(entity)
    saved: list[Path] = []

    for idx, img_tag in enumerate(soup.find_all("img", src=True)):
        if not isinstance(img_tag, Tag):
            continue
        src_raw = img_tag.get("src")
        src = src_raw.strip() if isinstance(src_raw, str) else ""
        if not src:
            continue

        # --- data URI ---------------------------------------------------------
        data_match = re.match(r"data:(image/[\w+.-]+);base64,(.+)", src, re.DOTALL)
        if data_match:
            mime = data_match.group(1)
            ext = mimetypes.guess_extension(mime) or ".bin"
            ext = ".jpg" if ext == ".jpe" else ext
            try:
                image_bytes = base64.b64decode(data_match.group(2))
                out_path = images_dir / f"image_{idx:03d}{ext}"
                out_path.write_bytes(image_bytes)
                saved.append(out_path)
            except Exception as e:
                logger.warning(f"[images/html] data-URI decode failed (idx {idx}): {e}")
            continue

        # --- remote URL -------------------------------------------------------
        try:
            full_url = urljoin(entity.url, src) if not src.startswith("http") else src
        except Exception:
            continue
        if not full_url.startswith("http"):
            continue

        if _is_private_url(full_url):
            logger.warning(
                f"[images/html] skipping non-public URL (SSRF guard): {full_url}"
            )
            continue

        try:
            r = _HTML_IMAGE_SESSION.get(full_url, timeout=15)
            r.raise_for_status()
            content_type = (
                r.headers.get("Content-Type", "image/jpeg").split(";")[0].strip()
            )
            ext = mimetypes.guess_extension(content_type) or ".jpg"
            ext = ".jpg" if ext == ".jpe" else ext
            out_path = images_dir / f"image_{idx:03d}{ext}"
            out_path.write_bytes(r.content)
            saved.append(out_path)
        except Exception as e:
            logger.warning(f"[images/html] download failed for {full_url}: {e}")

    logger.info(
        f"[images/html] extracted {len(saved)} image(s) from {entity.file_path}"
    )
    return saved


def extract_images_from_pdf(entity: EntityToClean) -> list[Path]:
    """
    Extract embedded images from a PDF using PyMuPDF (fitz).
    Each unique xref is saved once; duplicates (same image referenced on
    multiple pages) are skipped to avoid inflating the image set.
    """
    import fitz  # PyMuPDF — already a transitive dep via pymupdf4llm

    images_dir = _images_dir(entity)
    saved: list[Path] = []
    seen_xrefs: set[int] = set()

    doc = fitz.open(entity.file_path.as_posix())
    try:
        for page_num in range(len(doc)):
            for img_info in doc.get_page_images(page_num, full=True):
                xref = img_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    img_data = doc.extract_image(xref)
                    ext = f".{img_data['ext']}"
                    out_path = images_dir / f"image_{len(saved):03d}{ext}"
                    out_path.write_bytes(img_data["image"])
                    saved.append(out_path)
                except Exception as e:
                    logger.warning(f"[images/pdf] xref {xref} extraction failed: {e}")
    finally:
        doc.close()

    logger.info(f"[images/pdf] extracted {len(saved)} image(s) from {entity.file_path}")
    return saved


def extract_images_from_docx(entity: EntityToClean) -> list[Path]:
    """
    Extract embedded images from a DOCX file via python-docx relationship
    parts.  Only image relationships are processed; other media (audio, video)
    are skipped.
    """
    from docx import Document  # python-docx

    images_dir = _images_dir(entity)
    saved: list[Path] = []

    doc = Document(entity.file_path.as_posix())
    for idx, rel in enumerate(doc.part.rels.values()):
        if "image" not in rel.reltype:
            continue
        try:
            image_part = rel.target_part
            ext = Path(image_part.partname).suffix or ".png"
            out_path = images_dir / f"image_{idx:03d}{ext}"
            out_path.write_bytes(image_part.blob)
            saved.append(out_path)
        except Exception as e:
            logger.warning(f"[images/docx] rel {idx} extraction failed: {e}")

    logger.info(
        f"[images/docx] extracted {len(saved)} image(s) from {entity.file_path}"
    )
    return saved


def extract_images_from_pptx(entity: EntityToClean) -> list[Path]:
    """
    Extract embedded images from a PPTX file via python-pptx.
    Iterates every slide and picks out picture shapes (shape_type == 13,
    i.e. MSO_SHAPE_TYPE.PICTURE).
    """
    from pptx import Presentation  # python-pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.picture import Picture

    images_dir = _images_dir(entity)
    saved: list[Path] = []

    prs = Presentation(entity.file_path.as_posix())
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            if not isinstance(shape, Picture):
                continue
            try:
                image = shape.image
                ext = f".{image.ext}"
                out_path = images_dir / f"image_{len(saved):03d}{ext}"
                out_path.write_bytes(image.blob)
                saved.append(out_path)
            except Exception as e:
                logger.warning(f"[images/pptx] shape extraction failed: {e}")

    logger.info(
        f"[images/pptx] extracted {len(saved)} image(s) from {entity.file_path}"
    )
    return saved


def extract_images(entity: EntityToClean, file_type: str) -> list[Path]:
    """
    Dispatcher: call the right extractor for the given file type.
    Returns an empty list for formats with no image content (e.g. .txt).
    Any unexpected top-level failure is caught, logged, and treated as
    zero images so the rest of the pipeline is never blocked.
    """
    extractors = {
        ".html": extract_images_from_html,
        ".pdf": extract_images_from_pdf,
        ".docx": extract_images_from_docx,
        ".pptx": extract_images_from_pptx,
    }
    extractor = extractors.get(file_type)
    if extractor is None:
        return []
    try:
        return extractor(entity)
    except Exception as e:
        logger.error(
            f"[images] top-level extraction error for {entity.file_path} ({file_type}): {e}"
        )
        return []


# ---------------------------------------------------------------------------
# Logging setup
# ---------------------------------------------------------------------------


def set_up_logging(entity: EntityToClean) -> None:
    fmt = logging.Formatter(
        "[%(asctime)s] %(levelname)s [%(name)s.%(funcName)s:%(lineno)d] %(message)s"
    )
    root = logging.getLogger()
    root.setLevel(logging.INFO)

    # Attach stdout handler once only (root logger persists across processes).
    if not root.handlers:
        stdout_handler = logging.StreamHandler()
        stdout_handler.setFormatter(fmt)
        root.addHandler(stdout_handler)

    # Attach a per-job file handler, avoiding duplicates.
    # Resolve both paths to absolute strings before comparing so that relative
    # vs absolute differences do not create spurious duplicates.
    job_logger = logging.getLogger(__name__)
    log_path = entity.logs_path.resolve().as_posix()
    for handler in job_logger.handlers:
        if (
            isinstance(handler, logging.FileHandler)
            and Path(handler.baseFilename).resolve().as_posix() == log_path
        ):
            break
    else:
        file_handler = logging.FileHandler(log_path)
        file_handler.setFormatter(fmt)
        job_logger.addHandler(file_handler)


# ---------------------------------------------------------------------------
# Main single-file task
# ---------------------------------------------------------------------------


def clean_file_task(entity: EntityToClean) -> None:
    with catch_error(entity):
        set_up_logging(entity)
        logger.info(f"Cleaning file {entity.file_path.as_posix()}")

        ruuter_timeout = 30  # seconds

        with entity.meta_data_path.open("r") as f:
            metadata = json.load(f)

        file_type = metadata.get("file_type", "")

        # Only fetch Vault secrets when LLM is actually requested.
        # The service must start and process non-LLM jobs even when Vault
        # is not configured.
        client = None
        deployment_name = None
        if entity.use_llm:
            secrets = get_vault_secrets()
            client = _make_openai_client(secrets)
            deployment_name = secrets.azure_openai_deployment

        if file_type == ".html":
            cleaned_text = clean_html(entity, client, deployment_name)
            logger.info(f"Cleaned as HTML for {entity.file_path.as_posix()}")
        elif file_type == ".pdf":
            cleaned_text = clean_pdf(entity)
            logger.info(
                f"Cleaned as PDF (pymupdf4llm) for {entity.file_path.as_posix()}"
            )
        elif file_type in (".pptx", ".ppt"):
            cleaned_text = clean_any_file(entity)
            logger.info(
                f"Cleaned as PPTX (unstructured) for {entity.file_path.as_posix()}"
            )
        elif file_type in _PLAIN_TEXT_TYPES:
            cleaned_text = clean_plain_text(entity)
            logger.info(f"Cleaned as plain text for {entity.file_path.as_posix()}")
        else:
            cleaned_text = clean_any_file(entity)
            logger.info(
                f"Cleaned as unstructured file for {entity.file_path.as_posix()}"
            )

        # Normalise excessive whitespace and strip presentational cruft from
        # any embedded HTML tables produced by trafilatura/BS/LLM/Unstructured.
        # Plain-text sources are passed through verbatim: the file *is* the
        # text, so rewriting HTML the author wrote by hand is not ours to do.
        cleaned_text = normalize_newlines(cleaned_text)
        if file_type not in _PLAIN_TEXT_TYPES:
            cleaned_text = clean_html_tables(cleaned_text)

        # Detect language
        detected_language = None
        if cleaned_text.strip():
            try:
                detected_language = detect(cleaned_text)
                logger.info(
                    f"Detected language: {detected_language} for {entity.file_path.as_posix()}"
                )
            except LangDetectException as e:
                logger.error(
                    f"Language detection failed for {entity.file_path.as_posix()}: {e}"
                )

        # Write and upload cleaned text
        cleaned_text_filename = entity.directory_path / "cleaned.txt"
        with cleaned_text_filename.open("w") as f:
            f.write(cleaned_text)

        r = requests.post(
            f"{settings.ruuter_internal}/ckb/pipeline/upload-file-sync",
            json={"source_file_path": cleaned_text_filename.as_posix()},
            timeout=ruuter_timeout,
        )
        r.raise_for_status()
        try:
            response_json = r.json()
        except ValueError as exc:
            raise RuntimeError(
                "Invalid JSON response from Ruuter for cleaned text upload"
            ) from exc
        if not isinstance(response_json, dict) or "response" not in response_json:
            raise RuntimeError(
                "Missing 'response' key in Ruuter response for cleaned text upload"
            )
        uploaded_cleaned_text_url = response_json["response"]
        logger.info(f"Saved cleaned text for {entity.file_path.as_posix()}")

        # Write and upload cleaned metadata.
        # Both mutations happen before the write so the file is always consistent.
        # Remove any stale top-level "language" key the scrapper may have placed there;
        # the canonical location is metadata["metadata"]["language"].
        metadata.pop("language", None)
        metadata["metadata"]["cleaned"] = True
        metadata["metadata"]["language"] = detected_language
        cleaned_metadata_filename = entity.directory_path / "cleaned.meta.json"
        with cleaned_metadata_filename.open("w") as f:
            json.dump(metadata, f)

        r = requests.post(
            f"{settings.ruuter_internal}/ckb/pipeline/upload-file-sync",
            json={"source_file_path": cleaned_metadata_filename.as_posix()},
            timeout=ruuter_timeout,
        )
        r.raise_for_status()
        try:
            response_json = r.json()
        except ValueError as exc:
            raise RuntimeError(
                "Invalid JSON response from Ruuter for cleaned metadata upload"
            ) from exc
        if not isinstance(response_json, dict) or "response" not in response_json:
            raise RuntimeError(
                "Missing 'response' key in Ruuter response for cleaned metadata upload"
            )
        uploaded_cleaned_metadata_url = response_json["response"]
        logger.info(f"Saved cleaned metadata for {entity.file_path.as_posix()}")

        # Extract and upload images (only when explicitly requested)
        # Failures on individual images are logged but never raise — a missing
        # image must not abort an otherwise-successful cleaning job.
        uploaded_image_urls: list[str] = []
        extracted_images = (
            extract_images(entity, file_type) if entity.extract_images else []
        )
        for img_path in extracted_images:
            try:
                r = requests.post(
                    f"{settings.ruuter_internal}/ckb/pipeline/upload-file-sync",
                    json={"source_file_path": img_path.as_posix()},
                    timeout=ruuter_timeout,
                )
                r.raise_for_status()
                img_url = r.json().get("response", "")
                if img_url:
                    uploaded_image_urls.append(img_url)
                    logger.info(f"Uploaded image {img_path.name} -> {img_url}")
                else:
                    logger.warning(
                        f"[images] Ruuter returned empty URL for {img_path.name}"
                    )
            except Exception as e:
                logger.warning(f"[images] failed to upload {img_path.name}: {e}")

        if uploaded_image_urls:
            logger.info(
                f"Uploaded {len(uploaded_image_urls)}/{len(extracted_images)} image(s) "
                f"for {entity.file_path.as_posix()}"
            )

        # Update the database record
        # NOTE: Ruuter's update-cleaned-file endpoint must be extended to
        # accept the image_urls field so images are persisted in the DB.
        r = requests.post(
            f"{settings.ruuter_internal}/ckb/source-file/update-cleaned-file",
            json={
                "base_id": entity.source_file_id,
                "cleaned_data_url": uploaded_cleaned_text_url,
                "cleaned_metadata_url": uploaded_cleaned_metadata_url,
                "image_urls": uploaded_image_urls,
            },
            timeout=ruuter_timeout,
        )
        r.raise_for_status()

        # All uploads confirmed -- safe to delete the working directory
        cleanup_directory(entity)


# ---------------------------------------------------------------------------
# Batch source task
# ---------------------------------------------------------------------------


def _to_local_path(path_or_url: str | None) -> Path | None:
    if not path_or_url:
        return None
    return Path("/" + path_or_url.replace("uploads/", "").lstrip("/"))


def clean_source_task(task: SourceCleaningTask) -> None:
    logs_path = Path(task.logs_path)
    logs_path.parent.mkdir(parents=True, exist_ok=True)
    logs_path.touch(exist_ok=True)

    for file in task.files:
        try:
            requests.post(
                f"{settings.ruuter_internal}/ckb/source-file/update-scrapped-file-stop-scrapping",
                json={"base_id": file.base_id, "status": "cleaning"},
                timeout=30,
            )

            fallback_directory = (
                Path("/scrapped-data")
                / task.agency_base_id
                / file.source_base_id
                / file.base_id
            )
            file_path = (
                _to_local_path(file.original_data_url)
                or fallback_directory / "source.html"
            )
            meta_data_path = (
                _to_local_path(file.original_metadata_url)
                or fallback_directory / "source.meta.json"
            )

            entity = EntityToClean(
                file_path=file_path,
                meta_data_path=meta_data_path,
                directory_path=fallback_directory,
                source_file_id=file.base_id,
                url=file.url,
                logs_path=logs_path,
                source_base_id=file.source_base_id,
                agency_base_id=task.agency_base_id,
                source_run_report_base_id=task.source_run_report_base_id,
                use_llm=task.use_llm,
                use_llm_correction=task.use_llm_correction,
                extract_images=task.extract_images,
            )
            clean_file_task(entity)
        except Exception as e:
            # Catches both ValidationError and any runtime error -- both are
            # reported the same way and must not stop the remaining files.
            send_error(
                file.url,
                "cleaning",
                str(e),
                file.source_base_id,
                task.agency_base_id,
                task.source_run_report_base_id,
            )

    # Upload the cleaning log
    cleaning_log_url = ""
    if logs_path.exists():
        try:
            upload_result = requests.post(
                f"{settings.ruuter_internal}/ckb/pipeline/upload-file-sync",
                json={"source_file_path": logs_path.as_posix()},
                timeout=30,
            )
            upload_result.raise_for_status()
            cleaning_log_url = upload_result.json().get("response", "")
        except requests.exceptions.RequestException as e:
            logger.error(f"Failed to upload cleaning log file: {e}")

    # Update the run report
    try:
        response = requests.post(
            f"{settings.ruuter_internal}/ckb/reports/update",
            json={
                "baseId": task.source_run_report_base_id,
                "scrapingFinishedAt": datetime.datetime.now(datetime.UTC).isoformat(),
                "scrapingLogUrl": task.scraping_log_url,
                "cleaningLogUrl": cleaning_log_url,
            },
            timeout=30,
        )
        response.raise_for_status()
    except requests.exceptions.RequestException as e:
        logger.error(f"Failed to update report with cleaning log URL: {e}")

    # These two final status updates are best-effort -- wrap each independently
    # so a failure on the first does not prevent the second from running.
    try:
        requests.post(
            f"{settings.ruuter_internal}/ckb/source/update-status",
            json={"source_id": task.source_base_id, "status": "finished"},
            timeout=30,
        )
    except requests.exceptions.RequestException as e:
        logger.error(f"Failed to update source status: {e}")

    try:
        requests.post(
            f"{settings.ruuter_internal}/ckb/agency/update-zip-dirty",
            json={"sourceId": task.source_base_id, "agencyId": task.agency_base_id},
            timeout=30,
        )
    except requests.exceptions.RequestException as e:
        logger.error(f"Failed to update agency zip-dirty: {e}")

    try:
        requests.get(
            f"{settings.ruuter_internal}/ckb/pipeline/zip",
            timeout=300,
        )
    except requests.exceptions.RequestException as e:
        logger.error(f"Failed to trigger immediate zip after first-time cleaning: {e}")
